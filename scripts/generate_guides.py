import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_banner(doc, title_text, subtitle_text):
    # Banner title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(title_text)
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 76, 129) # Azul corporativo

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run(subtitle_text)
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 76, 129)

def add_subsection_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)

def add_body_paragraph(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix + " ")
        r_bold.font.name = "Arial"
        r_bold.font.size = Pt(10)
        r_bold.font.bold = True
        r_bold.font.color.rgb = RGBColor(30, 41, 59)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.italic = italic
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet_point(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix + " ")
        r_bold.font.name = "Arial"
        r_bold.font.size = Pt(10)
        r_bold.font.bold = True
        r_bold.font.color.rgb = RGBColor(30, 41, 59)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(51, 65, 85)

def add_callout_box(doc, text, title="NOTA IMPORTANTE", bg_hex="F0FDF4", border_hex="16A34A"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(10.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 76, 129)
    
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_admin_guide(output_path):
    doc = Document()
    
    # Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    add_header_banner(doc, "GUÍA DEL ADMINISTRADOR - GUÍA LOCAL", "Manual de Operaciones, Gestión Integral, Alta de Comercios y Facultades del Sistema")
    
    add_callout_box(doc, "Este documento detalla todas las herramientas, facultades y procesos que el Administrador de Guía Local tiene a disposición para gestionar comercios, alojamientos, actividades, facturación, banners, métricas y seguridad.", "RESUMEN EJECUTIVO")

    add_section_heading(doc, "1. Acceso y Seguridad al Panel Administrativo")
    add_body_paragraph(doc, "El acceso al panel central de administración está protegido mediante autenticación de administrador única.")
    add_bullet_point(doc, "https://[tu-dominio]/admin", "URL de Ingreso:")
    add_bullet_point(doc, "Usuario administrador y contraseña maestra asignada en el servidor.", "Credenciales:")
    add_bullet_point(doc, "Cierre de sesión seguro y protección contra accesos no autorizados en todas las rutas protegidas.", "Seguridad:")

    add_section_heading(doc, "2. Dashboard y Monitoreo en Tiempo Real")
    add_body_paragraph(doc, "Al ingresar a la consola administrativa, el panel presenta una visión consolidada del estado general de la plataforma:")
    add_bullet_point(doc, "Registro en tiempo real de accesos únicos por página y secciones turísticas.", "Visitas Globales (PageVisits):")
    add_bullet_point(doc, "Métricas acumuladas de usuarios que contactaron a comercios vía WhatsApp.", "Conversiones a WhatsApp:")
    add_bullet_point(doc, "Cantidad de aperturas del menú o catálogo digital de cada comercio.", "Aperturas de Menú:")
    add_bullet_point(doc, "Resumen inmediato de suscripciones al día vs. cuentas vencidas con botón de alerta.", "Estado de Cobranzas:")

    add_section_heading(doc, "3. Procedimiento para Dar de Alta un Comercio / Negocio")
    add_body_paragraph(doc, "El Administrador puede dar de alta un comercio en cuestión de minutos. Dependiendo de la categoría del negocio, se selecciona el módulo correspondiente:")
    
    add_subsection_heading(doc, "A. Comercios Gastronómicos y Generales (/admin/commerces o /admin/businesses)")
    add_bullet_point(doc, "Ingresar el nombre oficial del establecimiento (ej: 'Restaurante Las Lengas').", "1. Nombre:")
    add_bullet_point(doc, "Seleccionar el rubro correspondiente (Restaurante, Bar, Cervecería, Supermercado, Farmacia, etc.).", "2. Categoría / Rubro:")
    add_bullet_point(doc, "Cargar número con código de área e internacional, sin símbolos ni espacios (ej: 5492942123456). Esto habilitará el botón directo de WhatsApp para los turistas.", "3. Teléfono de WhatsApp:")
    add_bullet_point(doc, "Redactar una descripción seductora destacando especialidades, platos recomendados y comodidades.", "4. Descripción y Servicios:")
    add_bullet_point(doc, "Subir fotografía de portada y logo representativo (almacenamiento en alta velocidad Cloudinary / Vercel Blob).", "5. Imágenes:")
    add_bullet_point(doc, "Indicar coordenadas (Latitud y Longitud) para ubicar el negocio con pin interactivo en el mapa turístico.", "6. Geolocalización:")
    add_bullet_point(doc, "Establecer franja horaria (ej: '12:00 a 16:00 y 20:00 a 00:30') y activar el cronograma público.", "7. Horarios de Atención:")
    add_bullet_point(doc, "Activar el interruptor 'Verificado' para otorgar la insignia de confianza al comercio.", "8. Insignia Verificada:")

    add_subsection_heading(doc, "B. Alojamientos y Cabañas (/admin/accommodations)")
    add_body_paragraph(doc, "Permite registrar complejos turísticos, cabañas, hosterías y campings con características específicas (Wi-Fi, estacionamiento, vista al lago, parrilla, piscina, etc.).")

    add_subsection_heading(doc, "C. Aventuras, Excursiones y Guías (/admin/adventures)")
    add_body_paragraph(doc, "Módulo para paseos náuticos, trekking guiado, cabalgatas, pesca deportiva, kayak y deportes de nieve con desglose de servicios y requerimientos.")

    add_subsection_heading(doc, "D. Servicios Locales y Profesionales (/admin/local-services)")
    add_body_paragraph(doc, "Para rubros de mantenimiento, poda, raleo, auxilio mecánico, gomería, cerrajería, salud y servicios de emergencia.")

    add_section_heading(doc, "4. Creación y Asignación de Usuarios para el Comercio (Portal Comercial)")
    add_body_paragraph(doc, "Para que el dueño del comercio pueda autogestionar su menú, horarios y fotos sin requerir asistencia constante del Administrador:")
    add_bullet_point(doc, "El administrador crea la cuenta en PortalUser ingresando el correo electrónico del comerciante y una contraseña inicial.", "1. Generación de Acceso:")
    add_bullet_point(doc, "En la ficha del comercio, el administrador vincula el ID del comercio con el PortalUser creado.", "2. Vinculación del Comercio:")
    add_bullet_point(doc, "Se le entrega al comerciante el enlace de acceso (/portal-comercial/login) con sus credenciales para que gestione su negocio.", "3. Entrega de Credenciales:")

    add_section_heading(doc, "5. Módulo de Facturación, Planes y Cobranzas (/admin/billing)")
    add_body_paragraph(doc, "Control financiero centralizado de todas las suscripciones de los comercios:")
    add_bullet_point(doc, "Configurar si el cliente posee Plan Básico, Plan Completo o Plan Destacado.", "Planes:")
    add_bullet_point(doc, "Monitoreo con alerta visual (Verde: Al día / Rojo: Vencido).", "Vencimientos (dueDate):")
    add_bullet_point(doc, "Activar o desactivar posiciones de banners incluidos en su abono (Portada, Top, Medio, Inferior).", "Banners Incluidos:")
    add_bullet_point(doc, "Permite aplicar bonificaciones parciales o totales (ej: 'Mes de apertura bonificado', 'Descuento anual').", "Bonificaciones:")
    add_bullet_point(doc, "Registrar cada abono recibido, monto y fecha para conservar el historial contable.", "Historial de Pagos:")

    add_section_heading(doc, "6. Gestión de Publicidad y Banners (/admin/publicity)")
    add_body_paragraph(doc, "Administración de los espacios publicitarios destacados en la aplicación:")
    add_bullet_point(doc, "Inicio, Qué Comer, Dónde Dormir, Guía Local, Aventuras.", "Páginas destino:")
    add_bullet_point(doc, "Definir en qué sección (1 a 4) y qué orden de prioridad (1, 2, 3) se muestra el anuncio.", "Ubicación y Orden:")
    add_bullet_point(doc, "Programar fecha de inicio y vencimiento automático para que el banner se oculte al expirar el contrato.", "Vigencia:")
    add_bullet_point(doc, "Banner Grande (full width) o Banner Chico (tarjeta lateral o intercalada).", "Formatos:")

    add_section_heading(doc, "7. Moderación de Reseñas y Calificaciones (/admin/reviews)")
    add_body_paragraph(doc, "El Administrador tiene la potestad de auditar todas las opiniones dejadas por los visitantes:")
    add_bullet_point(doc, "Revisar comentarios antes o después de ser publicados.", "Aprobación:")
    add_bullet_point(doc, "Eliminar valoraciones que infrinjan las normas de convivencia o contengan spam.", "Eliminación:")

    add_section_heading(doc, "8. Notificaciones Push Masivas (/admin/notifications)")
    add_body_paragraph(doc, "Herramienta de alto impacto para fidelización de usuarios:")
    add_bullet_point(doc, "Enviar avisos de ofertas de fin de semana, eventos culturales, estado de rutas o emergencias a todos los usuarios que tengan instalada la aplicación (PWA).", "Envío Directo:")
    add_bullet_point(doc, "Consulta de suscriptores activos, envíos exitosos y porcentaje de clics.", "Estadísticas:")

    add_section_heading(doc, "9. Gestión de Contenido y Blog Turístico (/admin/blog)")
    add_body_paragraph(doc, "Publicación de artículos informativos, recomendaciones turísticas, senderos y atractivos para potenciar el posicionamiento SEO en Google y dar valor agregado al visitante.")

    add_section_heading(doc, "10. Marcadores del Mapa Interactivo (/admin/map)")
    add_body_paragraph(doc, "Creación y edición de puntos de interés clave: miradores, accesos a playas, cajeros automáticos, centros de salud y estaciones de servicio.")

    doc.save(output_path)
    print(f"Admin guide saved to: {output_path}")

def build_commerce_guide(output_path):
    doc = Document()
    
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    add_header_banner(doc, "GUÍA PARA EL COMERCIO - GUÍA LOCAL", "Manual de Registro, Gestión de Perfil, Carta Digital y Ventas por WhatsApp")
    
    add_callout_box(doc, "¡Te damos la bienvenida a Guía Local de Villa Pehuenia! Esta guía práctica te enseñará cómo dar de alta tu negocio, cómo ingresar a tu portal comercial y cómo aprovechar todas las herramientas para atraer más clientes y vender más.", "BIENVENIDO A GUÍA LOCAL")

    add_section_heading(doc, "1. ¿Qué es Guía Local y qué beneficios te brinda?")
    add_body_paragraph(doc, "Guía Local es la plataforma digital de referencia para turistas y residentes de Villa Pehuenia y Moquehue. Al tener tu comercio en la plataforma obtendrás:")
    add_bullet_point(doc, "Tu negocio disponible las 24 horas para miles de turistas que buscan qué comer, dónde alojarse o qué servicios contratar.", "Presencia Digital Permanente:")
    add_bullet_point(doc, "Los clientes tocan un botón y te escriben directamente a tu WhatsApp, sin comisiones ni intermediarios.", "Contacto Directo por WhatsApp:")
    add_bullet_point(doc, "Tus clientes ven tu menú, platos del día o catálogo de servicios desde su celular con fotos y precios actualizados.", "Carta Digital sin Costo de Impresión:")
    add_bullet_point(doc, "Aparición exacta en el mapa de Villa Pehuenia con botón para 'Cómo llegar'.", "Ubicación en el Mapa:")

    add_section_heading(doc, "2. ¿Cómo darte de alta en Guía Local?")
    add_body_paragraph(doc, "El proceso de alta es sumamente ágil y se realiza en 3 simples pasos:")
    
    add_subsection_heading(doc, "Paso 1: Solicitud de Alta con el Administrador")
    add_body_paragraph(doc, "Comunícate con el equipo de Guía Local y proporciona los datos principales de tu comercio:")
    add_bullet_point(doc, "Nombre de fantasía de tu comercio.", "Nombre:")
    add_bullet_point(doc, "Gastronomía, Alojamiento, Actividad / Aventura o Servicio Profesional.", "Rubro:")
    add_bullet_point(doc, "Número donde recibes pedidos y consultas de clientes.", "WhatsApp comercial:")
    add_bullet_point(doc, "Tu correo electrónico para crear tu cuenta personal de acceso.", "Email de contacto:")

    add_subsection_heading(doc, "Paso 2: Recepción de tus Credenciales de Acceso")
    add_body_paragraph(doc, "Una vez habilitado tu negocio en la plataforma, recibirás tu usuario y contraseña temporal para ingresar a tu panel de control exclusivo.")

    add_subsection_heading(doc, "Paso 3: Primer Ingreso a tu Portal Comercial")
    add_bullet_point(doc, "Ingresa desde tu celular o computadora a: https://[tu-dominio]/portal-comercial/login", "Acceso Web:")
    add_bullet_point(doc, "Escribe tu correo electrónico y la contraseña provista.", "Login:")
    add_bullet_point(doc, "¡Listo! Ya estás dentro de tu panel de autogestión.", "Ingreso:")

    add_section_heading(doc, "3. Cómo Gestionar el Perfil de tu Negocio (/portal-comercial/perfil)")
    add_body_paragraph(doc, "Dentro de la sección 'Mi Perfil' puedes configurar toda la información que verán los turistas:")
    add_bullet_point(doc, "Escribe un texto claro y atractivo sobre tu propuesta, tus especialidades gastronómicas, ambiente o servicios destacados.", "Descripción:")
    add_bullet_point(doc, "Verifica que tu número esté escrito en formato completo para que los clientes te contacten al instante.", "WhatsApp:")
    add_bullet_point(doc, "Indica tus días y horarios de atención. Puedes activar el indicador 'Abierto Ahora' para que los turistas sepan cuándo visitarte.", "Horarios de Atención:")
    add_bullet_point(doc, "Sube fotos luminosas de tu local, tus platos o tus cabañas. Las fotos de buena calidad duplican las consultas.", "Fotos y Logo:")
    add_bullet_point(doc, "Marca si cuentas con Wi-Fi, Estacionamiento, Medios de Pago Electrónico, Menú Celíaco/Vegano, Pet Friendly, etc.", "Servicios y Etiquetas:")

    add_section_heading(doc, "4. Cómo Administrar tu Carta o Catálogo Digital (/portal-comercial/menu)")
    add_body_paragraph(doc, "Si tienes un restaurante, bar, rotisería o negocio de productos, tu menú digital es tu mejor vendedor:")
    add_bullet_point(doc, "Haz clic en 'Nuevo Ítem', escribe el nombre del plato, agrega una descripción tentadora y ponle el precio.", "Agregar Producto:")
    add_bullet_point(doc, "Puedes cambiar los precios en cualquier momento desde tu celular. Los cambios se reflejan al instante en la guía.", "Actualización de Precios:")
    add_bullet_point(doc, "Si un ingrediente se agotó o el plato es de temporada, puedes desactivarlo temporalmente sin borrarlo.", "Pausar / Ocultar:")
    add_bullet_point(doc, "Sube fotos atractivas de tus platos estrella para estimular el apetito de los clientes.", "Fotos de Productos:")

    add_section_heading(doc, "5. Cómo Interpretar tus Estadísticas de Rendimiento (/portal-comercial)")
    add_body_paragraph(doc, "En la pantalla principal de tu Portal Comercial podrás evaluar el impacto real de tu negocio:")
    add_bullet_point(doc, "Cantidad de personas que abrieron tu ficha comercial este mes.", "👁️ Visitas Totales:")
    add_bullet_point(doc, "La métrica más valiosa: cuántos turistas presionaron el botón para enviarte un WhatsApp directo.", "📲 Clics a WhatsApp:")
    add_bullet_point(doc, "Cuántas veces los clientes consultaron tu carta o lista de precios.", "📋 Vistas de Menú:")

    add_section_heading(doc, "6. Consejos de Oro para Aumentar tus Ventas")
    add_bullet_point(doc, "Responde las consultas de WhatsApp en los primeros minutos. Los turistas en vacaciones eligen el primer lugar que les confirma disponibilidad o mesa.", "1. Rapidez en WhatsApp:")
    add_bullet_point(doc, "Si abres más temprano o cambias horario en temporada alta o baja, modifícalo en tu perfil para evitar confusiones.", "2. Horarios al Día:")
    add_bullet_point(doc, "Consulta con el Administrador para solicitar el sello oficial de 'Comercio Verificado' y destacar con banners en la portada de Guía Local.", "3. Destacarte con Banners:")

    add_section_heading(doc, "7. Soporte y Atención al Comerciante")
    add_body_paragraph(doc, "¿Tienes dudas para cargar tu menú o necesitas cambiar algún dato especial? Comunícate con el equipo de soporte de Guía Local a través de los canales oficiales habilitados.")

    doc.save(output_path)
    print(f"Commerce guide saved to: {output_path}")

if __name__ == "__main__":
    base_dir = r"c:\Users\gus16\Desktop\Antigravity\Proyectos\Guia Local"
    admin_path = os.path.join(base_dir, "Guia_Administrador_GuiaLocal.docx")
    commerce_path = os.path.join(base_dir, "Guia_Comercio_GuiaLocal.docx")
    
    build_admin_guide(admin_path)
    build_commerce_guide(commerce_path)
